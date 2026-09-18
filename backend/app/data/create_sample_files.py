import os
import sys

# Ensure folder exists
os.makedirs("sample_data/resumes", exist_ok=True)
os.makedirs("sample_data/job_descriptions", exist_ok=True)

try:
    import docx
except ImportError:
    print("python-docx is not installed. Exiting.")
    sys.exit(1)

def create_docx_resume(filename, name, email, phone, education, experience, skills_list, projects, certifications):
    doc = docx.Document()
    
    # Title/Header
    doc.add_heading(name, 0)
    doc.add_paragraph(f"Email: {email} | Phone: {phone}")
    
    # Skills Section
    doc.add_heading("Skills", level=1)
    doc.add_paragraph(", ".join(skills_list))
    
    # Education Section
    doc.add_heading("Education", level=1)
    doc.add_paragraph(education)
    
    # Experience Section
    doc.add_heading("Experience", level=1)
    doc.add_paragraph(experience)
    
    # Projects Section
    doc.add_heading("Projects", level=1)
    doc.add_paragraph(projects)
    
    # Certifications Section
    doc.add_heading("Certifications", level=1)
    doc.add_paragraph(certifications)
    
    # Save
    doc.save(os.path.join("sample_data/resumes", filename))
    print(f"Created resume: {filename}")

def main():
    # Candidate 1: ML Engineer
    create_docx_resume(
        filename="candidate1_ml.docx",
        name="Anirudh Saini",
        email="anirudh.saini@example.com",
        phone="+91-9876543210",
        education="B.Tech in AI & Data Science - GPA 3.9/4.0",
        experience="Machine Learning Engineer Intern at Tech Corp (1 year). Worked on building NLP classification engines, training statistical models, and processing datasets using Python and Scikit-learn.",
        skills_list=["Python", "Machine Learning", "NLP", "SQL", "TensorFlow", "Pandas", "NumPy"],
        projects="NLP Text Summarizer: Created a REST API with FastAPI that extracts features and summarizes files.",
        certifications="TensorFlow Developer Certificate, AWS Cloud Practitioner"
    )
    
    # Candidate 2: Java Developer
    create_docx_resume(
        filename="candidate2_java.docx",
        name="Rajesh Kumar",
        email="rajesh.kumar@example.com",
        phone="+91-8765432109",
        education="MCA (Master of Computer Applications) - 2024",
        experience="Software Developer at Systems Ltd (3 years). Focused on server-side enterprise apps using Java and Spring Boot. Managed database schemas and queries in PostgreSQL.",
        skills_list=["Java", "Spring Boot", "SQL", "PostgreSQL", "Git", "Docker"],
        projects="E-Commerce API: Built transaction management features using Spring Boot and Postgres.",
        certifications="Oracle Certified Associate Java SE Developer"
    )
    
    # Candidate 3: Full Stack React/Node Developer
    create_docx_resume(
        filename="candidate3_fullstack.docx",
        name="Priya Patel",
        email="priya.patel@example.com",
        phone="+91-7654321098",
        education="B.Tech in Computer Science - GPA 3.8/4.0",
        experience="Frontend Engineer at WebScale (2 years). Designed dashboard layouts using HTML/CSS, React.js, and JavaScript. Built microservices on the backend using Node.js.",
        skills_list=["JavaScript", "React", "Node.js", "HTML", "CSS", "Git", "TypeScript"],
        projects="Recruitment CRM: Created interactive graphs and applicant tables using React and Node.",
        certifications="Scrum Master Certification (PSM I)"
    )
    
    # Candidate 4: Data Analyst
    create_docx_resume(
        filename="candidate4_datascience.docx",
        name="Sanjay Sharma",
        email="sanjay.sharma@example.com",
        phone="+91-6543210987",
        education="B.Sc in Mathematics & Computer Science",
        experience="Data Analyst at Analytics Inc (3 years). Cleaned large datasets in Python using Pandas and NumPy. Generated executive reports and dashboards in Power BI and Tableau.",
        skills_list=["SQL", "Python", "Power BI", "Tableau", "Pandas", "NumPy", "Excel"],
        projects="Sales Performance Board: Consolidated multi-region DB records into interactive Power BI report.",
        certifications="Google Professional Data Engineer"
    )
    
    # Candidate 5: Cloud DevOps
    create_docx_resume(
        filename="candidate5_devops.docx",
        name="Neha Gupta",
        email="neha.gupta@example.com",
        phone="+91-5432109876",
        education="B.E. in Information Technology",
        experience="DevOps Systems Engineer (4 years). Configured CI/CD pipelines, containerized microservices in Docker, and managed cluster scale in Kubernetes. Orchestrated clouds in AWS.",
        skills_list=["Docker", "Kubernetes", "AWS", "Git", "Linux", "Python", "Azure"],
        projects="Automated Deployments: Set up Kubernetes clusters on AWS using Terraform templates.",
        certifications="AWS Certified Solutions Architect - Professional"
    )

    # Candidate 6: Frontend Developer & UI/UX Designer
    create_docx_resume(
        filename="candidate6_frontend.docx",
        name="Vikram Aditya",
        email="vikram.aditya@example.com",
        phone="+91-9988776655",
        education="Bachelor of Design (B.Des) - UI/UX & Web Interaction",
        experience="UI/UX Developer at Creative Studios (2 years). Designed modern dark-theme dashboard mockups in Figma, and implemented responsive frontend interfaces using HTML, CSS, React, and Sass.",
        skills_list=["React", "HTML", "CSS", "JavaScript", "Figma", "Sass", "Git"],
        projects="Interactive Design Hub: Coded reusable React layout kits and Figma tokens.",
        certifications="Google UX Design Professional Certificate"
    )

    # Candidate 7: Python Backend Engineer
    create_docx_resume(
        filename="candidate7_backend.docx",
        name="Aarav Mehta",
        email="aarav.mehta@example.com",
        phone="+91-8877665544",
        education="B.Tech in Information Technology",
        experience="Backend Engineer at DevOps Solutions (3 years). Built high-throughput API microservices using Python and FastAPI. Managed caching mechanisms in Redis and relational databases in PostgreSQL.",
        skills_list=["Python", "FastAPI", "PostgreSQL", "Redis", "Docker", "Git", "SQL"],
        projects="Microservices Gateway: Built high-throughput backend router in FastAPI with Redis caching.",
        certifications="Certified Professional Python Programmer"
    )

    # Candidate 8: Database Administrator (DBA)
    create_docx_resume(
        filename="candidate8_dba.docx",
        name="Meera Nair",
        email="meera.nair@example.com",
        phone="+91-7766554433",
        education="MCA in Systems & Network Engineering",
        experience="Senior Database Administrator at Financial Services Corp (5 years). Handled enterprise scale queries, database tuning, and high-availability setups across Oracle, PostgreSQL, and MySQL databases.",
        skills_list=["SQL", "PostgreSQL", "MySQL", "Oracle", "Database Tuning", "NoSQL", "Git"],
        projects="Data Warehouse Migration: Migrated legacy transactional DBs into unified clean Postgres schemas.",
        certifications="Oracle Database Administration Professional"
    )

    # Candidate 9: Cloud Solutions Architect
    create_docx_resume(
        filename="candidate9_cloud.docx",
        name="Rohit Verma",
        email="rohit.verma@example.com",
        phone="+91-6655443322",
        education="M.Tech in Cloud Computing - IIT",
        experience="Cloud Architect (6 years). Designed high-availability enterprise networks in AWS. Managed container orchestration in Kubernetes and infrastructure deployment scripting via Terraform templates.",
        skills_list=["AWS", "Docker", "Kubernetes", "Terraform", "Linux", "Python", "IAM", "Git"],
        projects="Global Scaling Plan: Orchestrated multi-region clouds deploying auto-scaling Kubernetes nodes.",
        certifications="AWS Certified Solutions Architect - Associate"
    )

    # Candidate 10: QA Automation Engineer
    create_docx_resume(
        filename="candidate10_qa.docx",
        name="Sneha Reddy",
        email="sneha.reddy@example.com",
        phone="+91-5544332211",
        education="B.Tech in Electronics & Communication",
        experience="QA Engineer at Automation Inc (3 years). Built robust automated test suites in Java using Selenium and Javascript using Cypress. Wired automated testing triggers into GitLab CI/CD pipelines.",
        skills_list=["Java", "Selenium", "Cypress", "QA", "Git", "Docker", "CI/CD"],
        projects="E2E Test Platform: Coded parallel Cypress test suite reducing regression times by 40%.",
        certifications="Certified ISTQB Software Tester"
    )
    
    # Create 3 job descriptions in text files
    jds = {
        "ml_engineer.txt": """
Job Title: Machine Learning Engineer
Experience Required: 2 years
Required Skills: Python, Machine Learning, NLP, SQL
Preferred Skills: TensorFlow, Docker
Description: We are seeking a Machine Learning Engineer to design and deploy high-performance NLP models. Candidate will preprocess text, extract features, and maintain REST endpoints.
        """,
        "fullstack_developer.txt": """
Job Title: Full Stack Developer
Experience Required: 3 years
Required Skills: JavaScript, React, Node.js, SQL, PostgreSQL
Preferred Skills: Git, FastAPI
Description: Join our product team as a Full Stack Developer. You will build frontend user dashboards in React and implement backend APIs in Node.js, integrated with PostgreSQL databases.
        """,
        "data_analyst.txt": """
Job Title: Data Analyst
Experience Required: 1 year
Required Skills: SQL, Python, Power BI, Tableau
Preferred Skills: Git, PostgreSQL
Description: Looking for a Data Analyst to write optimized SQL queries, inspect data quality in Pandas, and present key metrics to management using Power BI and Tableau.
        """
    }
    
    for filename, text in jds.items():
        with open(os.path.join("sample_data/job_descriptions", filename), "w") as f:
            f.write(text.strip())
        print(f"Created JD: {filename}")

if __name__ == "__main__":
    main()
